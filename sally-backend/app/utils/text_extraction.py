"""
Text extraction utilities for various file formats.
Centralized module to avoid code duplication.
"""

from pathlib import Path
from typing import Optional
import logging

logger = logging.getLogger(__name__)

# Check if text extraction libraries are available
try:
    from pypdf import PdfReader
    from docx import Document
    import openpyxl
    TEXT_EXTRACTION_AVAILABLE = True
except ImportError:
    TEXT_EXTRACTION_AVAILABLE = False
    logger.warning("Text extraction libraries not available. Install: pypdf, python-docx, openpyxl")


def extract_text_from_pdf(file_path: Path) -> str:
    """
    Extract text from PDF file.
    
    Args:
        file_path: Path to the PDF file
        
    Returns:
        Extracted text or error message
    """
    if not TEXT_EXTRACTION_AVAILABLE:
        return "Text extraction libraries not available"

    try:
        reader = PdfReader(file_path)
        text = ""
        for page in reader.pages:
            text += page.extract_text() + "\n"
        return text.strip()
    except Exception as e:
        logger.error(f"Error extracting PDF text from {file_path}: {e}")
        return f"Error extracting PDF text: {str(e)}"


def extract_text_from_docx(file_path: Path) -> str:
    """
    Extract text from DOCX file.
    
    Args:
        file_path: Path to the DOCX file
        
    Returns:
        Extracted text or error message
    """
    if not TEXT_EXTRACTION_AVAILABLE:
        return "Text extraction libraries not available"

    try:
        doc = Document(file_path)
        text = "\n".join([paragraph.text for paragraph in doc.paragraphs])
        return text.strip()
    except Exception as e:
        logger.error(f"Error extracting DOCX text from {file_path}: {e}")
        return f"Error extracting DOCX text: {str(e)}"


def extract_text_from_xlsx(file_path: Path) -> str:
    """
    Extract text from XLSX/Excel file.
    
    Args:
        file_path: Path to the XLSX file
        
    Returns:
        Extracted text or error message
    """
    if not TEXT_EXTRACTION_AVAILABLE:
        return "Text extraction libraries not available"

    try:
        workbook = openpyxl.load_workbook(file_path)
        text = ""
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            text += f"Sheet: {sheet_name}\n"
            for row in sheet.iter_rows(values_only=True):
                row_text = "\t".join(str(cell) for cell in row if cell is not None)
                if row_text.strip():
                    text += row_text + "\n"
            text += "\n"
        return text.strip()
    except Exception as e:
        logger.error(f"Error extracting Excel text from {file_path}: {e}")
        return f"Error extracting Excel text: {str(e)}"


# Alias for backward compatibility
extract_text_from_excel = extract_text_from_xlsx


def extract_text_from_csv(file_path: Path) -> str:
    """
    Extract text from CSV file.
    
    Args:
        file_path: Path to the CSV file
        
    Returns:
        Extracted text or error message
    """
    try:
        with open(file_path, 'r', encoding='utf-8') as f:
            return f.read()
    except UnicodeDecodeError:
        try:
            with open(file_path, 'r', encoding='latin-1') as f:
                return f.read()
        except Exception as e:
            logger.error(f"Error extracting CSV text from {file_path}: {e}")
            return f"Error extracting CSV text: {str(e)}"
    except Exception as e:
        logger.error(f"Error extracting CSV text from {file_path}: {e}")
        return f"Error extracting CSV text: {str(e)}"


def extract_text_from_txt(file_path: Path, encoding: str = 'utf-8') -> str:
    """
    Extract text from TXT file.
    
    Args:
        file_path: Path to the TXT file
        encoding: File encoding (default: utf-8)
        
    Returns:
        Extracted text or error message
    """
    try:
        with open(file_path, 'r', encoding=encoding) as f:
            return f.read().strip()
    except Exception as e:
        logger.error(f"Error reading TXT file {file_path}: {e}")
        return f"Error reading TXT file: {str(e)}"


def extract_text_from_file(file_path: Path, file_extension: Optional[str] = None) -> str:
    """
    Extract text from file based on extension.
    
    Args:
        file_path: Path to the file
        file_extension: File extension (optional, will be detected from path)
        
    Returns:
        Extracted text or error message
    """
    if file_extension is None:
        file_extension = file_path.suffix.lower()
    
    extractors = {
        '.pdf': extract_text_from_pdf,
        '.docx': extract_text_from_docx,
        '.doc': extract_text_from_docx,  # Try docx extractor for .doc
        '.xlsx': extract_text_from_xlsx,
        '.xls': extract_text_from_xlsx,  # Try xlsx extractor for .xls
        '.csv': extract_text_from_csv,
        '.txt': extract_text_from_txt,
    }
    
    extractor = extractors.get(file_extension)
    if extractor:
        return extractor(file_path)
    else:
        return f"Unsupported file format: {file_extension}"

